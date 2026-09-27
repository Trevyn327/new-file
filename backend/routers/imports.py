import io
import csv
from fastapi import APIRouter, Depends, UploadFile, File, HTTPException
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from database import products, categories, sales, purchases, customers, suppliers, now_iso, log_activity
from auth import get_current_user, new_id

router = APIRouter(prefix='/api', tags=['io'])

PRODUCT_HEADERS = ['sku', 'name', 'category', 'unit', 'purchase_price', 'selling_price', 'stock', 'low_stock_threshold']
XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def _csv(rows, header, filename):
    buf = io.StringIO(); w = csv.writer(buf); w.writerow(header); w.writerows(rows); buf.seek(0)
    return StreamingResponse(iter([buf.getvalue()]), media_type='text/csv',
                             headers={'Content-Disposition': f'attachment; filename={filename}'})


def _xlsx(rows, header, filename):
    wb = Workbook(); ws = wb.active; ws.append(header)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return StreamingResponse(buf, media_type=XLSX_MIME,
                             headers={'Content-Disposition': f'attachment; filename={filename}'})


@router.get('/products/template.csv')
async def product_template(user=Depends(get_current_user)):
    sample = [['RICE-5KG', 'Basmati Rice 5kg', 'Grocery', 'bag', '420', '520', '40', '15']]
    return _csv(sample, PRODUCT_HEADERS, 'product_template.csv')


@router.get('/export/products.xlsx')
async def export_products_xlsx(user=Depends(get_current_user)):
    cat = {c['id']: c['name'] async for c in categories.find()}
    rows = []
    async for p in products.find().sort('name', 1):
        rows.append([p.get('sku'), p['name'], cat.get(p.get('category_id'), ''), p.get('unit'),
                     p.get('purchase_price'), p.get('selling_price'), p.get('stock'), p.get('low_stock_threshold')])
    return _xlsx(rows, PRODUCT_HEADERS, 'products.xlsx')


@router.get('/export/sales.xlsx')
async def export_sales_xlsx(user=Depends(get_current_user)):
    cust = {c['id']: c['name'] async for c in customers.find()}
    rows = []
    async for s in sales.find().sort('date', -1):
        rows.append([s['invoice_no'], s['date'], cust.get(s.get('customer_id'), 'Walk-in'),
                     s.get('total'), s.get('amount_paid'), s.get('balance_due'), s.get('status')])
    return _xlsx(rows, ['Invoice', 'Date', 'Customer', 'Total', 'Paid', 'Balance', 'Status'], 'sales.xlsx')


@router.get('/export/purchases.xlsx')
async def export_purchases_xlsx(user=Depends(get_current_user)):
    sup = {s['id']: s['name'] async for s in suppliers.find()}
    rows = []
    async for p in purchases.find().sort('date', -1):
        rows.append([p['date'], sup.get(p.get('supplier_id'), 'Cash Supplier'), p.get('total'),
                     p.get('amount_paid'), p.get('balance_due'), p.get('status')])
    return _xlsx(rows, ['Date', 'Supplier', 'Total', 'Paid', 'Balance', 'Status'], 'purchases.xlsx')


@router.post('/products/import')
async def import_products(file: UploadFile = File(...), user=Depends(get_current_user)):
    content = await file.read()
    fname = (file.filename or '').lower()
    records = []
    try:
        if fname.endswith('.xlsx'):
            wb = load_workbook(io.BytesIO(content)); ws = wb.active
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise ValueError('empty')
            header = [str(h).strip().lower() if h is not None else '' for h in rows[0]]
            for r in rows[1:]:
                records.append({header[i]: r[i] for i in range(len(header))})
        else:
            text = content.decode('utf-8-sig')
            for r in csv.DictReader(io.StringIO(text)):
                records.append({(k or '').strip().lower(): v for k, v in r.items()})
    except Exception:
        raise HTTPException(status_code=400, detail='Could not read the file. Please use the template format.')

    cat_by_name = {c['name'].lower(): c['id'] async for c in categories.find()}

    def num(v, d=0):
        try:
            return float(v)
        except Exception:
            return d

    added = updated = 0
    failed = []
    for idx, row in enumerate(records, start=2):
        try:
            nm = (str(row.get('name')) if row.get('name') else '').strip()
            if not nm:
                failed.append({'row': idx, 'reason': 'Missing name'}); continue
            sku = (str(row.get('sku')) if row.get('sku') else '').strip() or nm[:6].upper().replace(' ', '')
            catname = (str(row.get('category')) if row.get('category') else '').strip()
            cat_id = None
            if catname:
                cid = cat_by_name.get(catname.lower())
                if not cid:
                    cid = new_id()
                    await categories.insert_one({'id': cid, 'name': catname, 'created_at': now_iso()})
                    cat_by_name[catname.lower()] = cid
                cat_id = cid
            existing = await products.find_one({'sku': sku})
            if existing:
                await products.update_one({'id': existing['id']}, {'$set': {
                    'name': nm, 'category_id': cat_id or existing.get('category_id'),
                    'unit': (row.get('unit') or existing.get('unit') or 'pcs'),
                    'purchase_price': num(row.get('purchase_price'), existing.get('purchase_price', 0)),
                    'selling_price': num(row.get('selling_price'), existing.get('selling_price', 0)),
                    'low_stock_threshold': num(row.get('low_stock_threshold'), existing.get('low_stock_threshold', 5))}})
                updated += 1
            else:
                pp = num(row.get('purchase_price'))
                await products.insert_one({'id': new_id(), 'sku': sku, 'name': nm, 'category_id': cat_id,
                    'barcode': '', 'unit': (row.get('unit') or 'pcs'), 'purchase_price': pp,
                    'selling_price': num(row.get('selling_price')), 'stock': num(row.get('stock')),
                    'avg_cost': pp, 'image_url': '', 'low_stock_threshold': num(row.get('low_stock_threshold'), 5),
                    'created_at': now_iso()})
                added += 1
        except Exception as e:
            failed.append({'row': idx, 'reason': str(e)[:80]})
    await log_activity(user, 'import', 'product', None, {'added': added, 'updated': updated, 'failed': len(failed)})
    return {'added': added, 'updated': updated, 'failed': failed}
